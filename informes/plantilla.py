"""Lin-AI · plantilla de correo y resumen automático de adjuntos.

El correo lo monta siempre el código con esta plantilla; la IA no escribe el correo.
De la IA solo se acepta, como mucho:
- un título corto para el asunto (el código añade el prefijo y la fecha), y
- una nota breve en texto plano, que pasa por `limpiar_nota` (sin HTML, sin enlaces,
  sin direcciones y sin cifras que no salgan de los adjuntos).

El resumen de cada adjunto (filas, columnas, totales, apartados) lo calcula el código
leyendo el propio archivo, así que no gasta tokens y las cifras son las reales.
"""
import csv
import html
import re
import unicodedata
from datetime import datetime
from pathlib import Path

PREFIJO_ASUNTO = "[Lin-AI SPSIL]"
MAX_TITULO = 90
MAX_NOTA = 600
MAX_LINEAS_NOTA = 6
MAX_FILAS_LEIDAS = 100_000
MAX_TOTALES = 4

COLOR = {"marca": "#1F4E78", "fondo": "#F3F5F8", "borde": "#DCE3EC", "texto": "#1E2A36",
         "suave": "#5F6B7A", "aviso": "#FFF4D6", "aviso_borde": "#E8C468"}
ETIQUETA = {"xlsx": ("EXCEL", "#1D6F42"), "xlsm": ("EXCEL", "#1D6F42"), "csv": ("CSV", "#1D6F42"),
            "docx": ("WORD", "#2B579A"), "pdf": ("PDF", "#B3261E")}


class NotaNoValida(ValueError):
    """La nota de la IA no cumple las reglas; el mensaje dice qué quitar."""


# ---------------------------------------------------------------- números

def fmt_num(v):
    """Formato español: 1.234,56 (sin decimales si es entero)."""
    if isinstance(v, bool) or not isinstance(v, (int, float)):
        return str(v)
    if float(v).is_integer():
        return f"{int(v):,}".replace(",", ".")
    return f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def fmt_tamano(n):
    for unidad in ("B", "KB", "MB"):
        if n < 1024 or unidad == "MB":
            return f"{n:.0f} {unidad}" if unidad == "B" else f"{n:.1f} {unidad}".replace(".", ",")
        n /= 1024


def _es_numero(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _totales(columnas, filas):
    """Suma de las columnas que son numéricas en al menos el 80 % de las filas."""
    totales = []
    for i, col in enumerate(columnas):
        valores = [f[i] for f in filas if i < len(f) and f[i] not in (None, "")]
        numeros = [v for v in valores if _es_numero(v)]
        if numeros and len(numeros) >= len(valores) * 0.8 and not _parece_codigo(col):
            totales.append({"columna": col, "suma": round(sum(numeros), 2)})
    return totales[:MAX_TOTALES]


def _parece_codigo(col):
    """Columnas que son números pero no tiene sentido sumar (códigos, años, meses...)."""
    partes = re.findall(r"[a-z]+", _sin_tildes(str(col)).lower())
    return any(p.startswith(("cod", "num", "telef")) or p in ("id", "ano", "anio", "mes", "dia", "cp", "nif", "ejercicio")
               for p in partes)


def _sin_tildes(t):
    return "".join(ch for ch in unicodedata.normalize("NFD", t) if unicodedata.category(ch) != "Mn")


# ---------------------------------------------------------------- resumen de adjuntos

def _resumen_xlsx(ruta):
    from openpyxl import load_workbook
    libro = load_workbook(ruta, read_only=True, data_only=True)
    try:
        hoja = libro["Datos"] if "Datos" in libro.sheetnames else libro.active
        filas = []
        for n, fila in enumerate(hoja.iter_rows(values_only=True)):
            if n > MAX_FILAS_LEIDAS:
                break
            filas.append(list(fila))
        titulo = filas[0][0] if filas and isinstance(filas[0][0], str) else ""
        # Cabecera: primera fila con al menos 2 textos y nada más que textos
        cab = next((i for i, f in enumerate(filas)
                    if sum(bool(isinstance(v, str) and v.strip()) for v in f) >= 2
                    and all(v is None or isinstance(v, str) for v in f)), None)
        if cab is None:
            return {"titulo": titulo, "hojas": len(libro.sheetnames)}
        columnas = [str(v) for v in filas[cab] if v not in (None, "")]
        datos = [f[:len(columnas)] for f in filas[cab + 1:] if any(v not in (None, "") for v in f)]
        return {"titulo": titulo, "filas": len(datos), "columnas": columnas,
                "totales": _totales(columnas, datos), "hojas": len(libro.sheetnames)}
    finally:
        libro.close()


def _resumen_docx(ruta):
    from docx import Document
    doc = Document(ruta)
    titulo, apartados = "", []
    for p in doc.paragraphs:
        estilo = (p.style.name or "").lower()
        if not p.text.strip():
            continue
        if estilo == "title" and not titulo:
            titulo = p.text.strip()
        elif estilo.startswith("heading"):
            apartados.append(p.text.strip())
    filas_tablas = sum(max(len(t.rows) - 1, 0) for t in doc.tables)
    palabras = sum(len(p.text.split()) for p in doc.paragraphs)
    return {"titulo": titulo, "apartados": apartados[:8], "tablas": len(doc.tables),
            "filas": filas_tablas or None, "palabras": palabras}


def _resumen_csv(ruta):
    with open(ruta, encoding="utf-8-sig", errors="replace", newline="") as f:
        muestra = f.read(4096)
        f.seek(0)
        try:
            dialecto = csv.Sniffer().sniff(muestra, delimiters=";,\t")
        except csv.Error:
            dialecto = csv.excel
        lector = csv.reader(f, dialecto)
        columnas = next(lector, [])
        filas = sum(1 for _ in lector)
    return {"filas": filas, "columnas": columnas}


def resumir_adjunto(ruta):
    """Ficha de un adjunto calculada por código. Nunca falla: si no se puede leer, da lo básico."""
    ruta = Path(ruta)
    ext = ruta.suffix.lower().lstrip(".")
    ficha = {"nombre": ruta.name, "tipo": ext, "tamano": fmt_tamano(ruta.stat().st_size)}
    lector = {"xlsx": _resumen_xlsx, "xlsm": _resumen_xlsx, "docx": _resumen_docx, "csv": _resumen_csv}.get(ext)
    if lector:
        try:
            ficha.update({k: v for k, v in lector(ruta).items() if v not in (None, "", [])})
        except Exception:  # archivo dañado o protegido: se envía igual, sin detalle
            ficha["aviso"] = "No se pudo leer el contenido para el resumen."
    return ficha


def lineas_ficha(ficha):
    """Frases cortas del resumen de un adjunto (las usan el HTML y el texto plano)."""
    lineas = []
    if ficha.get("filas") is not None and ficha.get("tipo") != "docx":
        lineas.append(f"{fmt_num(ficha['filas'])} filas de datos")
    if ficha.get("columnas"):
        cols = ficha["columnas"]
        extra = f" y {len(cols) - 6} más" if len(cols) > 6 else ""
        lineas.append(f"{len(cols)} columnas: {', '.join(cols[:6])}{extra}")
    for t in ficha.get("totales") or []:
        lineas.append(f"Total {t['columna']}: {fmt_num(t['suma'])}")
    if ficha.get("apartados"):
        lineas.append("Apartados: " + ", ".join(ficha["apartados"]))
    if ficha.get("tablas"):
        filas = f" ({fmt_num(ficha['filas'])} filas)" if ficha.get("filas") else ""
        lineas.append(f"{ficha['tablas']} tabla{'s' if ficha['tablas'] > 1 else ''}{filas}")
    if ficha.get("aviso"):
        lineas.append(ficha["aviso"])
    return lineas


# ---------------------------------------------------------------- lo que aporta la IA

_RE_NUMERO = re.compile(r"\d[\d.,]*\d|\d")


def _clave_numero(texto):
    """'1.234,50' -> '123450' (solo los dígitos) para comparar."""
    return re.sub(r"\D", "", texto).lstrip("0") or "0"


def cifras_permitidas(fichas):
    permitidas = set()
    for f in fichas:
        for v in [f.get("filas"), len(f.get("columnas") or []), f.get("tablas")] + \
                 [t["suma"] for t in f.get("totales") or []]:
            if _es_numero(v):
                permitidas |= {_clave_numero(fmt_num(v)), _clave_numero(str(v))}
        for texto in [f.get("titulo") or "", f["nombre"]] + list(f.get("apartados") or []):
            permitidas |= {_clave_numero(n) for n in _RE_NUMERO.findall(str(texto))}
    return permitidas


def limpiar_titulo(titulo, fichas=()):
    """Título del asunto: una línea, sin HTML ni símbolos raros. Si falta, sale del primer adjunto."""
    t = re.sub(r"<[^>]*>", "", str(titulo or ""))
    t = re.sub(r"[^\w\s\-–.,:/()%&ºª]", "", t, flags=re.UNICODE)
    t = " ".join(t.split())[:MAX_TITULO].strip(" -–.,:")
    if not t:
        t = next((f.get("titulo") for f in fichas if f.get("titulo")), "") or \
            (Path(fichas[0]["nombre"]).stem.rsplit("_", 2)[0].replace("_", " ") if fichas else "Informe")
    return t[:MAX_TITULO]


def asunto(titulo, fecha=None):
    return f"{PREFIJO_ASUNTO} {titulo} · {(fecha or datetime.now()):%d/%m/%Y}"


def limpiar_nota(nota, fichas=()):
    """Nota breve de la IA en texto plano. Lanza NotaNoValida si trae algo que no debe."""
    t = str(nota or "").replace("\r", "")
    t = re.sub(r"<[^>]*>", "", t)                                # sin HTML
    t = re.sub(r"[*_`#>|]+", "", t)                              # sin markdown
    if re.search(r"https?://|www\.", t, re.I):
        raise NotaNoValida("La nota no puede llevar enlaces.")
    if re.search(r"[\w.+-]+@[\w-]+\.[\w.]+", t):
        raise NotaNoValida("La nota no puede llevar direcciones de correo.")
    lineas = [" ".join(l.split()) for l in t.split("\n")]
    t = "\n".join(l for l in lineas if l)
    if len(t) > MAX_NOTA or t.count("\n") >= MAX_LINEAS_NOTA:
        raise NotaNoValida(f"La nota es demasiado larga (máximo {MAX_NOTA} caracteres y "
                           f"{MAX_LINEAS_NOTA} líneas). El resumen de los adjuntos ya lo pone el programa.")
    permitidas = cifras_permitidas(fichas)
    raras = [n for n in _RE_NUMERO.findall(t)
             if _clave_numero(n) not in permitidas and not _fecha_o_pequeno(n)]
    if raras:
        raise NotaNoValida("La nota tiene cifras que no salen de los adjuntos: "
                           f"{', '.join(dict.fromkeys(raras))}. Quítalas: las cifras las pone el programa.")
    return t


def _fecha_o_pequeno(n):
    """Se permiten días/meses (1-31) y años (2000-2100), p. ej. 'septiembre de 2026'."""
    if not n.isdigit():
        return False
    v = int(n)
    return 1 <= v <= 31 or 2000 <= v <= 2100


# ---------------------------------------------------------------- correo final

def componer(titulo, nota, fichas, autor="", modo_prueba=False, fecha=None):
    """Devuelve (texto_plano, html) del correo. Todo lo que viene de fuera se escapa."""
    fecha = fecha or datetime.now()
    e = html.escape
    n = len(fichas)
    intro = (f"Te envío {n} archivo{'s' if n != 1 else ''} generado{'s' if n != 1 else ''} por Lin-AI."
             if n else "Te escribo desde Lin-AI.")

    # ---- texto plano (para clientes que no muestran HTML)
    plano = ["Hola,", "", intro]
    if nota:
        plano += ["", nota]
    if fichas:
        plano += ["", "Qué va adjunto:"]
        for f in fichas:
            plano.append(f"- {f['nombre']} ({f['tamano']})")
            plano += [f"    {l}" for l in lineas_ficha(f)]
    plano += ["", "Las cifras de este resumen salen directamente de los archivos adjuntos.",
              f"Generado automáticamente por Lin-AI el {fecha:%d/%m/%Y a las %H:%M}"
              + (f" a petición de {autor}." if autor and autor != "anónimo" else ".")]
    if modo_prueba:
        plano.insert(0, "[CORREO DE PRUEBA: no se ha enviado]\n")

    # ---- HTML (tablas y estilos en línea: es lo que entienden Outlook y Gmail)
    c = COLOR
    tarjetas = []
    for f in fichas:
        etiqueta, color = ETIQUETA.get(f["tipo"], (f["tipo"].upper() or "ARCHIVO", c["suave"]))
        detalle = "".join(
            f'<tr><td style="padding:2px 0;font-size:13px;color:{c["texto"]};">• {e(l)}</td></tr>'
            for l in lineas_ficha(f))
        subtitulo = (f'<div style="font-size:13px;color:{c["suave"]};margin-top:2px;">{e(f["titulo"])}</div>'
                     if f.get("titulo") else "")
        tarjetas.append(f"""
<tr><td style="padding:0 0 12px 0;">
 <table role="presentation" width="100%" cellpadding="0" cellspacing="0"
        style="border:1px solid {c['borde']};border-radius:8px;background:#ffffff;">
  <tr><td style="padding:14px 16px;">
   <table role="presentation" width="100%" cellpadding="0" cellspacing="0"><tr>
    <td width="64" valign="top">
     <div style="background:{color};color:#ffffff;font:bold 11px Arial,sans-serif;letter-spacing:.5px;
                 text-align:center;padding:6px 0;border-radius:4px;width:56px;">{e(etiqueta)}</div></td>
    <td valign="top" style="padding-left:10px;">
     <div style="font:bold 14px Arial,sans-serif;color:{c['texto']};word-break:break-all;">{e(f['nombre'])}</div>
     {subtitulo}
     <div style="font-size:12px;color:{c['suave']};margin-top:2px;">{e(f['tamano'])}</div>
    </td></tr></table>
   <table role="presentation" cellpadding="0" cellspacing="0" style="margin-top:8px;font-family:Arial,sans-serif;">
    {detalle}</table>
  </td></tr></table>
</td></tr>""")

    aviso = (f"""<tr><td style="padding:10px 24px;background:{c['aviso']};border-bottom:1px solid {c['aviso_borde']};
                 font:13px Arial,sans-serif;color:{c['texto']};">Correo de prueba: se ha guardado, no se ha enviado.</td></tr>"""
             if modo_prueba else "")
    nota_html = (f"""<p style="margin:0 0 16px 0;padding:12px 14px;background:{c['fondo']};border-left:3px solid {c['marca']};
                    border-radius:4px;">{e(nota).replace(chr(10), '<br>')}</p>""" if nota else "")
    adjuntos_html = (f"""<h2 style="font:bold 15px Arial,sans-serif;color:{c['marca']};margin:8px 0 12px 0;">
                    Qué va adjunto</h2>
                   <table role="presentation" width="100%" cellpadding="0" cellspacing="0">{''.join(tarjetas)}</table>"""
                     if fichas else "")
    pie_autor = f" a petición de {e(autor)}" if autor and autor != "anónimo" else ""

    cuerpo_html = f"""<!doctype html>
<html lang="es"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width">
<title>{e(titulo)}</title></head>
<body style="margin:0;padding:0;background:{c['fondo']};">
<table role="presentation" width="100%" cellpadding="0" cellspacing="0" style="background:{c['fondo']};">
<tr><td align="center" style="padding:24px 12px;">
 <table role="presentation" width="600" cellpadding="0" cellspacing="0"
        style="max-width:600px;width:100%;background:#ffffff;border-radius:10px;overflow:hidden;
               border:1px solid {c['borde']};">
  <tr><td style="background:{c['marca']};padding:18px 24px;">
   <div style="font:bold 12px Arial,sans-serif;letter-spacing:1.5px;color:#BFD3E6;">SPSIL · LIN-AI</div>
   <div style="font:bold 20px Arial,sans-serif;color:#ffffff;margin-top:4px;">{e(titulo)}</div>
   <div style="font:13px Arial,sans-serif;color:#D6E2EE;margin-top:4px;">{fecha:%d/%m/%Y}</div>
  </td></tr>
  {aviso}
  <tr><td style="padding:22px 24px 8px 24px;font:14px/1.5 Arial,sans-serif;color:{c['texto']};">
   <p style="margin:0 0 12px 0;">Hola,</p>
   <p style="margin:0 0 16px 0;">{e(intro)}</p>
   {nota_html}
   {adjuntos_html}
  </td></tr>
  <tr><td style="padding:14px 24px 20px 24px;border-top:1px solid {c['borde']};
                 font:12px/1.5 Arial,sans-serif;color:{c['suave']};">
   Las cifras de este resumen salen directamente de los archivos adjuntos.<br>
   Generado automáticamente por Lin-AI el {fecha:%d/%m/%Y a las %H:%M}{pie_autor}.
  </td></tr>
 </table>
</td></tr></table>
</body></html>"""
    return "\n".join(plano), cuerpo_html
